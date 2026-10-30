"""Build the current interview manuscript; optional authoring dependency: python-docx."""
from pathlib import Path
import argparse
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
p=argparse.ArgumentParser();p.add_argument('--output-dir',type=Path,default=ROOT/'output/submission');a=p.parse_args()
a.output_dir.mkdir(parents=True,exist_ok=True)
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name,size in [('Normal',10.5),('Title',25),('Heading 1',18),('Heading 2',12),('Caption',8.5)]:
 s=doc.styles[name];s.font.name='Arial Unicode MS';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
 s.paragraph_format.line_spacing=1.24;s.paragraph_format.space_after=Pt(8)
 if name.startswith('Heading'):s.paragraph_format.space_before=Pt(4)
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
doc.core_properties.title='RocoCoach 基于 LLM 的智能陪练教练'
doc.core_properties.subject='设计 实现 难点与验证'
doc.core_properties.author='RocoCoach';doc.core_properties.last_modified_by='RocoCoach'

def table(lines):
 rows=[[x.strip() for x in l.strip().strip('|').split('|')] for l in lines if '---' not in l]
 t=doc.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=([2,7.5,7.5] if rows[0][0]=='角色' else [2,4.5,10.5] if rows[0][0]=='层' else [7.5,4.75,4.75] if rows[0][0].startswith('相同') else [4,5,8])
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,txt in enumerate(row):
   c=cells[j];c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar);b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF3' if i==0 else ('F6F8FA' if i%2==0 else 'FFFFFF'));pr.append(sh)
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.18
   if j==0 and rows[0][0] in {'角色','层'}:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   r=p.add_run(txt);r.font.size=Pt(9.5);r.bold=i==0
  no=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(no)
  if i==0:
   repeat=OxmlElement('w:tblHeader');t.rows[-1]._tr.get_or_add_trPr().append(repeat)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
lines=(ROOT/'docs/SUBMISSION_DRAFT.md').read_text().splitlines();i=0;section_count=0
while i<len(lines):
 line=lines[i]
 if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):
  section_count+=1
  if section_count>1:
   if section_count==2:
    p=doc.add_paragraph();p.add_run().add_picture(str(ROOT/'docs/semantic-agent-execution-preview.png'),width=Cm(17))
    doc.add_paragraph('图 1 真实模型在隔离合成局理解上下文确认后实际开启进阶训练','Caption')
   doc.add_page_break()
  doc.add_paragraph(line[3:],'Heading 1')
 elif line.startswith('|'):
  group=[]
  while i<len(lines) and lines[i].startswith('|'):group.append(lines[i]);i+=1
  table(group);continue
 elif line.strip():
  p=doc.add_paragraph(line)
  if line.startswith('面试提交初稿'):p.runs[0].font.size=Pt(10)
 i+=1
for root in [doc._element,doc.styles.element]:
 for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
out=a.output_dir/'RocoCoach_设计与实现_面试初稿.docx';doc.save(out);print(out)
